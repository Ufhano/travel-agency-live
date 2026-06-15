#!/usr/bin/env python3
"""
Per-sheet .xlsx extraction for travel/finance tables.

Workbook → worksheets. Each sheet may have preamble rows (logos, titles, blanks);
we score rows to find a likely *header* row, then read tabular data below until
a long run of empty rows. Output is JSON: merged row dicts with _sheet and _sourceRow.

Requires: pip install -r requirements.txt (openpyxl, pandas optional for future).
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

MAX_SCAN_ROWS = 150
MAX_DATA_ROWS = 8000
EMPTY_STREAK_STOP = 30
HEADER_KEYWORDS = re.compile(
    r"name|client|email|phone|trip|tour|dest|amount|paid|balance|"
    r"invoice|ref|date|price|package|method|outstanding|total|due|pax|guest",
    re.I,
)


def cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return str(v)
    return str(v).strip()


def row_values(ws: Worksheet, row_idx: int, max_col: int) -> list[Any]:
    out: list[Any] = []
    for c in range(1, max_col + 1):
        out.append(ws.cell(row=row_idx, column=c).value)
    return out


def non_empty_count(vals: list[Any]) -> int:
    return sum(1 for v in vals if cell_str(v) != "")


def header_score(vals: list[str]) -> float:
    texts = [v for v in vals if v]
    if len(texts) < 2:
        return 0.0
    score = float(len(texts))
    for t in texts:
        tl = t.lower()
        if HEADER_KEYWORDS.search(tl):
            score += 4.0
        if len(t) <= 48 and re.search(r"[a-zA-Z]", t):
            score += 0.5
        if re.fullmatch(r"[\d\s,.$R€£ZARzar-]+", t) and len(t) > 4:
            score -= 2.5
    return score


def sheet_max_col(ws: Worksheet, scan_rows: int) -> int:
    mc = ws.max_column or 1
    mc = min(mc, 80)
    if mc < 2:
        mc = 12
    return mc


def best_header_row(ws: Worksheet, max_col: int, scan: int) -> int | None:
    best_r: int | None = None
    best_s = -1.0
    last = min(scan, ws.max_row or scan)
    for r in range(1, last + 1):
        raw = row_values(ws, r, max_col)
        vals = [cell_str(x) for x in raw]
        if non_empty_count(raw) < 2:
            continue
        s = header_score(vals)
        if s > best_s:
            best_s = s
            best_r = r
    if best_r is None or best_s < 3.5:
        return None
    return best_r


def sanitize_key(h: str, used: dict[str, int], col_idx: int) -> str:
    h = re.sub(r"\s+", " ", (h or "").strip())
    if not h:
        h = f"Column_{col_idx}"
    base = h[:120]
    key = base
    n = used.get(key, 0)
    used[key] = n + 1
    if n:
        key = f"{base}__{n}"
    return key


def extract_sheet(ws: Worksheet, sheet_name: str) -> dict[str, Any]:
    st = getattr(ws, "sheet_state", None)
    if st == "veryHidden":
        return {"name": sheet_name, "skipped": True, "reason": "veryHidden", "rows": []}

    max_col = sheet_max_col(ws, MAX_SCAN_ROWS)
    scan = min(MAX_SCAN_ROWS, ws.max_row or MAX_SCAN_ROWS)
    if scan < 2:
        return {"name": sheet_name, "rows": [], "header_row": None, "note": "empty sheet"}

    hdr = best_header_row(ws, max_col, scan)
    if hdr is None:
        return {
            "name": sheet_name,
            "rows": [],
            "header_row": None,
            "note": "no confident header row",
        }

    headers_raw = [cell_str(x) for x in row_values(ws, hdr, max_col)]
    used: dict[str, int] = {}
    headers = [sanitize_key(h, used, j) for j, h in enumerate(headers_raw)]

    rows_out: list[dict[str, Any]] = []
    empty_run = 0
    last_row = min(ws.max_row or hdr, hdr + MAX_DATA_ROWS)

    for r in range(hdr + 1, last_row + 1):
        raw = row_values(ws, r, max_col)
        if non_empty_count(raw) == 0:
            empty_run += 1
            if empty_run >= EMPTY_STREAK_STOP:
                break
            continue
        empty_run = 0

        obj: dict[str, Any] = {"_sheet": sheet_name, "_sourceRow": r}
        for j, key in enumerate(headers):
            if j >= len(raw):
                break
            v = raw[j]
            if v is not None and cell_str(v) != "":
                obj[key] = v
        data_keys = [k for k in obj if not k.startswith("_")]
        if not data_keys:
            continue
        rows_out.append(obj)

    return {
        "name": sheet_name,
        "header_row": hdr,
        "rows": rows_out,
        "column_headers": headers_raw,
    }


def is_ignored_sheet(name: str) -> bool:
    n = name.strip().lower()
    if not n:
        return False
    if n == "msc cruises 2026 admin":
        return True
    return bool(re.search(r"\sadmin\s*$", n))


def main() -> None:
    logs: list[str] = []
    if len(sys.argv) < 2:
        print(json.dumps({"ok": False, "error": "usage: excel_extract.py <file.xlsx>"}))
        sys.exit(2)

    path = Path(sys.argv[1])
    if not path.exists():
        print(json.dumps({"ok": False, "error": f"not found: {path}"}))
        sys.exit(2)

    if path.suffix.lower() != ".xlsx":
        print(
            json.dumps(
                {
                    "ok": False,
                    "error": "Python path supports .xlsx only; use Node for .xls/.csv.",
                }
            )
        )
        sys.exit(2)

    try:
        wb = load_workbook(path, read_only=False, data_only=True)
    except Exception as e:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": str(e), "logs": logs}))
        sys.exit(1)

    sheets_out: list[dict[str, Any]] = []
    merged: list[dict[str, Any]] = []

    try:
        for name in wb.sheetnames:
            if is_ignored_sheet(name):
                logs.append(f'Sheet "{name}": skipped (admin tab)')
                continue
            ws = wb[name]
            try:
                block = extract_sheet(ws, name)
            except Exception as e:  # noqa: BLE001
                logs.append(f'sheet "{name}" failed: {e}')
                continue
            sheets_out.append(block)
            merged.extend(block.get("rows") or [])
            hr = block.get("header_row")
            n = len(block.get("rows") or [])
            logs.append(f'Sheet "{name}": header_row={hr}, data_rows={n}')
    finally:
        wb.close()

    print(
        json.dumps(
            {
                "ok": True,
                "sheets": sheets_out,
                "rows": merged,
                "logs": logs,
            },
            default=str,
        )
    )


if __name__ == "__main__":
    main()
